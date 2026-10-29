"""Inspect downloaded attachments locally without executing their contents."""

from __future__ import annotations

import hashlib
import mimetypes
from pathlib import Path
import re
import warnings
import xml.etree.ElementTree as ET
import zipfile


_MAX_FILE_BYTES = 64 * 1024 * 1024
_MAX_XML_BYTES = 16 * 1024 * 1024
_MAX_PDF_PAGES = 200
_MAX_SHEET_ROWS = 10000
_MAX_SHEET_COLUMNS = 1024
_TEXT_SUFFIXES = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".jsonl", ".log",
    ".xml", ".html", ".htm", ".yaml", ".yml", ".srt", ".vtt",
}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
_MIME_TYPES = {
    ".hwpx": "application/hwp+zip",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class _Text:
    def __init__(self, limit: int):
        self.limit = limit
        self.value = ""
        self.truncated = False

    def add(self, value: str) -> None:
        remaining = self.limit - len(self.value)
        self.value += value[:remaining]
        self.truncated |= len(value) > remaining


def _error(result: dict, code: str, message: str) -> dict:
    result["status"] = "error"
    result["error"] = {"code": code, "message": message}
    return result


def _decode(data: bytes) -> tuple[str, str]:
    encodings = ("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "cp949")
    for encoding in encodings:
        try:
            value = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        if any(ord(char) < 32 and char not in "\r\n\t\f" for char in value):
            raise ValueError("The text attachment contains binary control bytes.")
        return value, encoding
    raise ValueError("Text is not valid UTF-8, BOM-marked UTF-16, or Korean CP949.")


def _xml_chunks(node):
    tag = node.tag.rsplit("}", 1)[-1]
    if tag == "t":
        yield "".join(node.itertext())
        return
    if tag == "tab":
        yield "\t"
    elif tag in {"br", "lineBreak"}:
        yield "\n"
    for child in node:
        yield from _xml_chunks(child)
    if tag == "p":
        yield "\n"


def _document(path: Path, result: dict, text: _Text) -> None:
    with zipfile.ZipFile(path) as archive:
        if path.suffix.lower() == ".docx":
            main = "word/document.xml"
            names = [main] + sorted(
                name for name in archive.namelist()
                if re.fullmatch(r"word/(?:header\d*|footer\d*|footnotes|endnotes)\.xml", name)
            )
            result["coverage"] = "document_body_headers_footnotes_text_only"
        else:
            names = sorted(
                (name for name in archive.namelist() if re.fullmatch(r"Contents/section\d+\.xml", name)),
                key=lambda name: int(re.search(r"(\d+)\.xml$", name)[1]),
            )
            if not names:
                raise ValueError("HWPX contains no Contents/sectionN.xml parts.")
            result["coverage"] = "hwpx_section_text_only"
        result["parts_extracted"] = []
        consumed = 0
        for name in names:
            size = archive.getinfo(name).file_size
            consumed += size
            if size > _MAX_XML_BYTES or consumed > _MAX_FILE_BYTES:
                text.truncated = True
                result["warnings"].append(f"XML extraction size limit reached at {name}.")
                break
            data = archive.read(name)
            if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", data.replace(b"\0", b""), re.IGNORECASE):
                raise ValueError("XML declarations with DTDs or entities are not accepted.")
            root = ET.fromstring(data)
            if result["parts_extracted"]:
                text.add("\n")
            for chunk in _xml_chunks(root):
                text.add(chunk)
                if text.truncated:
                    break
            result["parts_extracted"].append(name)
            if text.truncated:
                break


def _image(path: Path, result: dict) -> None:
    try:
        from PIL import Image
    except ImportError:
        result.update(supported=False, status="unsupported")
        result["warnings"].append("Install kakaocli-win[analysis] for image verification.")
        return
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(path) as picture:
            result.update(width=picture.width, height=picture.height, format=picture.format)
            result["frames"] = getattr(picture, "n_frames", 1)
            result["mime_type"] = Image.MIME.get(picture.format, result["mime_type"])
            picture.verify()
    result.update(
        coverage="verified_image_metadata_only",
        visual_paths=[str(path)],
        needs_visual_review=True,
    )
    result["warnings"].append("Image contents require visual review; no OCR or semantic analysis was performed.")


def _pdf_previews(path: Path, output_directory, result: dict) -> None:
    if output_directory is None:
        return
    try:
        import fitz
    except ImportError:
        result["warnings"].append("Install kakaocli-win[analysis] for PDF page previews.")
        return
    folder = Path(output_directory).expanduser().resolve() / f"{path.stem}-{result['sha256'][:12]}"
    folder.mkdir(parents=True, exist_ok=True)
    try:
        with fitz.open(path) as document:
            if document.needs_pass:
                result["warnings"].append("Encrypted PDF previews require a password.")
                return
            for index in range(min(3, document.page_count)):
                page = document[index]
                longest = max(page.rect.width, page.rect.height)
                scale = min(1.5, 1600 / longest) if longest > 0 else 1
                destination = folder / f"page-{index + 1}.png"
                page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).save(destination)
                result["visual_paths"].append(str(destination))
    except Exception as exc:
        result["warnings"].append(f"PDF preview failed: {type(exc).__name__}: {exc}")


def _pdf(path: Path, result: dict, text: _Text, output_directory) -> None:
    try:
        from pypdf import PdfReader
    except ImportError:
        result.update(supported=False, status="unsupported")
        result["warnings"].append("Install kakaocli-win[analysis] for PDF text extraction.")
        _pdf_previews(path, output_directory, result)
        return
    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        _error(result, "encrypted_pdf", "PDF requires a password; no text was extracted.")
        return
    result.update(page_count=len(reader.pages), pages_extracted=0, coverage="pdf_text_only")
    for index, page in enumerate(reader.pages):
        if index >= _MAX_PDF_PAGES:
            text.truncated = True
            result["warnings"].append("PDF text extraction stopped after 200 pages.")
            break
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            result["warnings"].append(f"Page {index + 1} extraction failed: {type(exc).__name__}: {exc}")
            result["status"] = "partial"
            continue
        text.add(f"[Page {index + 1}]\n{page_text}\n" if page_text else "")
        result["pages_extracted"] += 1
        if text.truncated:
            break
    if not text.value:
        result["warnings"].append("PDF contains no extracted text; scan images require visual review or OCR.")
    _pdf_previews(path, output_directory, result)
    result["needs_visual_review"] = True
    if result["visual_paths"]:
        result["coverage"] = "pdf_text_and_first_three_page_previews"


def _xlsx(path: Path, result: dict, text: _Text) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        result.update(supported=False, status="unsupported")
        result["warnings"].append("Install kakaocli-win[analysis] for XLSX cell extraction.")
        return
    # Bound decompressed workbook parts before handing the archive to openpyxl.
    with zipfile.ZipFile(path) as archive:
        xml_size = sum(info.file_size for info in archive.infolist() if info.filename.endswith(".xml"))
        if xml_size > _MAX_FILE_BYTES:
            raise ValueError("Workbook decompressed XML exceeds the extraction size limit.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    result.update(coverage="xlsx_cells_and_formulas_only", sheets=[])
    result["warnings"].append("Formulas are shown as stored and are never evaluated; charts and images are not extracted.")
    rows_read = 0
    limited_columns = False
    try:
        for sheet in workbook.worksheets:
            text.add(f"[Sheet: {sheet.title}]\n")
            result["sheets"].append(sheet.title)
            if (sheet.max_column or 0) > _MAX_SHEET_COLUMNS:
                result["status"] = "partial"
                limited_columns = True
                result["warnings"].append(f"Sheet {sheet.title} is limited to its first 1024 columns.")
            for row in sheet.iter_rows(values_only=True, max_col=min(sheet.max_column or _MAX_SHEET_COLUMNS, _MAX_SHEET_COLUMNS)):
                if rows_read >= _MAX_SHEET_ROWS:
                    text.truncated = True
                    result["warnings"].append("Workbook extraction stopped after 10000 rows.")
                    break
                rows_read += 1
                values = ["" if value is None else str(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if values:
                    text.add("\t".join(values) + "\n")
                if text.truncated:
                    break
            if text.truncated:
                break
    finally:
        workbook.close()
    text.truncated |= limited_columns
    result["rows_read"] = rows_read


def analyze_attachment(path, *, output_directory=None, max_chars: int = 60000) -> dict:
    """Return bounded local extraction and visual-review paths for an attachment.

    This function does not infer image contents, run OCR, execute macros or
    formulas, fetch embedded links, upload data, or invoke a model. PDF preview
    files are only created when an output directory is explicitly supplied.
    """
    if not isinstance(max_chars, int) or isinstance(max_chars, bool) or not 1 <= max_chars <= 1000000:
        raise ValueError("max_chars must be an integer between 1 and 1000000")
    local_path = Path(path).expanduser().resolve()
    suffix = local_path.suffix.lower()
    kind = (
        "image" if suffix in _IMAGE_SUFFIXES else
        "text" if suffix in _TEXT_SUFFIXES else
        suffix[1:] if suffix in {".pdf", ".docx", ".hwpx", ".xlsx"} else "unknown"
    )
    result = {
        "path": str(local_path), "name": local_path.name,
        "mime_type": _MIME_TYPES.get(suffix) or mimetypes.guess_type(local_path.name)[0] or "application/octet-stream",
        "kind": kind, "size_bytes": None, "sha256": None,
        "supported": kind != "unknown", "status": "ok", "text": "", "text_chars": 0,
        "truncated": False, "coverage": "none", "warnings": [], "visual_paths": [],
    }
    if not local_path.is_file():
        return _error(result, "file_not_found", "Attachment path does not point to an existing regular file.")
    text = _Text(max_chars)
    try:
        result["size_bytes"] = local_path.stat().st_size
        digest = hashlib.sha256()
        with local_path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
        result["sha256"] = digest.hexdigest()
        if not result["supported"]:
            result["status"] = "unsupported"
            result["warnings"].append("Supported formats: images, text/CSV/JSON, PDF, DOCX, HWPX, and XLSX.")
        elif result["size_bytes"] > _MAX_FILE_BYTES:
            _error(result, "file_too_large", "Attachment exceeds the 64 MiB local extraction limit.")
        elif kind == "text":
            decoded, encoding = _decode(local_path.read_bytes())
            text.add(decoded)
            result.update(encoding=encoding, coverage="text_content")
        elif kind == "image":
            _image(local_path, result)
        elif kind == "pdf":
            _pdf(local_path, result, text, output_directory)
        elif kind in {"docx", "hwpx"}:
            _document(local_path, result, text)
        elif kind == "xlsx":
            _xlsx(local_path, result, text)
    except Exception as exc:
        # Attachment corruption is data, rather than a traceback that hides all
        # successful earlier extraction from the CLI caller.
        if text.value:
            result["status"] = "partial"
            result["error"] = {"code": "extraction_failed", "message": f"{type(exc).__name__}: {exc}"}
        else:
            _error(result, "extraction_failed", f"{type(exc).__name__}: {exc}")
    result.update(text=text.value, text_chars=len(text.value), truncated=text.truncated)
    if text.truncated and result["status"] == "ok":
        result["status"] = "partial"
    return result
